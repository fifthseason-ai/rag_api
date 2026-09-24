"""Antonio's retrieval path at the PRODUCER (PACKET-1, C2P1): the new per-unit locators reach
the wire through the exact /query routes Core calls, are never fabricated, never cross a file or
entity boundary, and the SAME question asked twice returns a BYTE-IDENTICAL ordered result.

Richard's goal, verbatim: "design self checks and verify that antonio can access it and if asked
the same question 2 times, still surfaces the same answer". This file is the rag_api half of it.
Core's consumer (readPassageLocator) and the live Antonio journey are NOT proven here.

Core's request shape, read from release/richard-vibe 091d9729
(api/server/services/Knowledge/retrieval.js, api/app/clients/tools/util/fileSearch.js):
  knowledge leg  POST /query/{knowledgeId}  {query, k: legK, args: knowledge.args ?? {cohere: false}}
  file leg       POST /query                {file_id, query, k: legK}      (one call PER file)
  agent file     POST /query                {file_id, query, k: legK, entity_id: agentId}
  headers        Authorization: Bearer <KSPT-01 token: id, tid, ent, act=['read'], iat, exp>
                 ent = [userId (+ agentId for agent files)] on file legs; the knowledge ids on knowledge legs
  legK           45 for the Antonio file_search tools: max(DEFAULT_QUERY_K 15 * 3, MIN_RETRIEVAL_K 40).
                 /api/knowledge/search uses k = limit (default 10), a SMALLER k; the k-cut test below
                 covers a k smaller than the scope. Core's knowledge path segment is a knowledge-base
                 id; here it is the owning entity's id. rag_api takes the same code path for both: the
                 path id must be in `ent` and becomes the user_id row filter.

Retrieval paths, and HOW EACH TEST KNOWS WHICH ONE RAN (recorded per request and asserted):
  default  The app's configured defaults: HYBRID_SEARCH_ENABLED and RERANK_ENABLED both default to
           True in app/config.py. The keyword leg is LIVE: document_tsv is added the way TEMPO
           migration 10081 shapes it, as test_both_legs_expected_passage_and_version does. The
           Bedrock rerank CALL is replaced by a counting raise, so the real rerank() fallback runs,
           which is what production does when the Bedrock rerank errors.
           Asserted per request:
             * the keyword leg ran once and returned the matching rows;
             * a Bedrock rerank was attempted once on a file leg (no args);
             * no rerank was attempted on a knowledge leg sent with {cohere: false}.
           A SUCCESSFUL live Cohere reorder is an external boundary this suite cannot reach.
  dense    HYBRID_SEARCH_ENABLED=False and RERANK_ENABLED=False: the vector store alone.
           Asserted: the dense arm ran once, the keyword leg never ran, no rerank was attempted.

A keyword-SOURCED case stubs the dense arm to []. It proves the locators also survive the keyword
leg's own row-to-Document path. The default path never puts that path on the wire while the dense
arm returns every chunk, because RRF keeps the first-seen Document, which is the dense one.

NOT COVERED, on purpose (see the C2P1 FILES finding): EXACT-TIE ordering. There is no tie-break
anywhere: the dense leg orders by `distance` only and the keyword leg by `score` only. Rows at the
same distance or score therefore come back in heap/plan order, which is stable on a serial plan and
NOT on a parallel one. The repeat tests below ASSERT their no-tie premise (pairwise-distinct dense
distances and keyword scores), so a fixture change fails as a named precondition instead of flaking.
Fused (RRF) scores may tie. That is not checked, because RRF breaks such ties deterministically
(stable sort, dense-first insertion).

SYNTHETIC fixtures only, built at test time. Needs a Postgres with pgvector: RAG_TEST_PG_DSN
selects it and RAG_TEST_PG_REQUIRED=1 turns a missing DSN into an ERROR (CI sets both); with
neither set every test here SKIPS. Reuses the existing real-store harness (needs_pg, _real_store,
_client from test_parse_is_not_index; _DetEmb, _sqlalchemy_dsn from
test_empty_entitlement_query_path) -- no second harness.

FORWARD-COMPATIBLE BY DESIGN: no assertion pins the exact key set of a hit, of its document
object or of its metadata. An additive producer field, such as a per-hit score_kind, cannot turn
these tests red. Whole-body comparisons only ever compare two calls of the same code.
"""

import datetime
import io
import os
import zipfile

import jwt
import psycopg2
import pytest
from fastapi.testclient import TestClient

from tests.utils.test_docx_reading_order import DOCX_MIME, W_NS
from tests.utils.test_empty_entitlement_query_path import _DetEmb, _sqlalchemy_dsn
from tests.utils.test_parse_is_not_index import PG_DSN, _client, _real_store, needs_pg

_SECRET = "testsecret"
TENANT = "tenantA"
USER_A = "userA"
USER_B = "userB"
AGENT = "agentA"
CORE_K = 45
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MD_MIME = "text/markdown"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

DOCX_TOKEN = "C2P1QDOCXTOKEN"
HDR_TOKEN = "C2P1QHEADERTOKEN"
FTR_TOKEN = "C2P1QFOOTERTOKEN"
XLSX_TOKEN = "C2P1QXLSXTOKEN"
MD_TOKEN = "C2P1QMDTOKEN"
PRE_TOKEN = "C2P1QPREAMBLETOKEN"

F_DOCX, F_XLSX, F_MD = "c2p1-q-docx", "c2p1-q-xlsx", "c2p1-q-md"
F_AGENT_DOCX = "c2p1-q-docx-agent"
F_B_MD = "c2p1-q-md-b"

Q_REPEAT = "What is the quarterly revenue figure in euros?"


# --------------------------------------------------------------------------- SYNTHETIC builders
def docx_with_header_footer_SYNTHETIC():
    """Three body paragraphs (block_index 0,1,2; the answer is 1) plus a header and a footer
    part, whose units must carry NO block_index."""
    ct = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
        '<Override PartName="/word/header1.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.header+xml"/>'
        '<Override PartName="/word/footer1.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footer+xml"/>'
        "</Types>"
    )
    rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
        "</Relationships>"
    )
    body = (
        "<w:p><w:r><w:t>Introduction paragraph about general company matters.</w:t></w:r></w:p>"
        f"<w:p><w:r><w:t>The {DOCX_TOKEN} quarterly revenue figure is 4200000 euros.</w:t></w:r></w:p>"
        "<w:p><w:r><w:t>Closing remarks, disclaimers and appendix pointers.</w:t></w:r></w:p>"
    )
    document = (
        f'<?xml version="1.0"?><w:document xmlns:w="{W_NS}" xmlns:r="{R_NS}"><w:body>{body}'
        '<w:sectPr><w:headerReference w:type="default" r:id="rIdH"/>'
        '<w:footerReference w:type="default" r:id="rIdF"/></w:sectPr>'
        "</w:body></w:document>"
    )
    header = f'<?xml version="1.0"?><w:hdr xmlns:w="{W_NS}"><w:p><w:r><w:t>{HDR_TOKEN} confidential banner</w:t></w:r></w:p></w:hdr>'
    footer = f'<?xml version="1.0"?><w:ftr xmlns:w="{W_NS}"><w:p><w:r><w:t>{FTR_TOKEN} page footer</w:t></w:r></w:p></w:ftr>'
    drels = (
        '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rIdH" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/header" Target="header1.xml"/>'
        '<Relationship Id="rIdF" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footer" Target="footer1.xml"/>'
        "</Relationships>"
    )
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", ct)
        z.writestr("_rels/.rels", rels)
        z.writestr("word/document.xml", document)
        z.writestr("word/_rels/document.xml.rels", drels)
        z.writestr("word/header1.xml", header)
        z.writestr("word/footer1.xml", footer)
    return buf.getvalue()


def xlsx_with_blank_sheet_SYNTHETIC():
    """Three sheets:
    * Revenue: a header row and two data rows, occupied extent A1:C3.
    * Costs: a second NON-EMPTY sheet, extent A1:B2. It makes the per-sheet qualification of
      cell_range load-bearing.
    * Blank: no cell at all."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Revenue"
    for row in (["Region", "Metric", "Value"], ["EMEA", XLSX_TOKEN, 4200000], ["AMER", "revenue", 3100000]):
        ws.append(row)
    costs = wb.create_sheet("Costs")
    for row in (["Item", "Cost"], ["Rent", 100]):
        costs.append(row)
    wb.create_sheet("Blank")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def md_with_preamble_SYNTHETIC():
    return (
        f"{PRE_TOKEN} text before any heading.\n\n"
        "# Overview\n\nGeneral introduction to the report.\n\n"
        f"## Revenue\n\nThe {MD_TOKEN} quarterly revenue figure is 4200000 euros.\n\n"
        "## Costs\n\nMiscellaneous cost commentary.\n"
    ).encode("utf-8")


FIXTURES = {
    "docx": (F_DOCX, "report.docx", DOCX_MIME, docx_with_header_footer_SYNTHETIC),
    "xlsx": (F_XLSX, "book.xlsx", XLSX_MIME, xlsx_with_blank_sheet_SYNTHETIC),
    "md": (F_MD, "notes.md", MD_MIME, md_with_preamble_SYNTHETIC),
}


# --------------------------------------------------------------------------- harness glue
def _token(ent=(USER_A,), act=("read",), secret=_SECRET, tid=TENANT, hours=1):
    """Core's KSPT-01 claims: id, tid, ent, act, plus iat and exp (jsonwebtoken adds iat)."""
    os.environ["JWT_SECRET"] = _SECRET
    now = datetime.datetime.now(datetime.timezone.utc)
    payload = {
        "id": ent[0] if ent else "nobody", "tid": tid, "ent": list(ent), "act": list(act),
        "iat": now, "exp": now + datetime.timedelta(hours=hours),
    }
    return jwt.encode(payload, secret, algorithm="HS256")


def _core_hdr(ent=(USER_A,)):
    """Core's header: Bearer KSPT-01 token (read-only act) + JSON content type."""
    return {"Authorization": f"Bearer {_token(ent)}", "Content-Type": "application/json"}


def _write_hdr(ent=(USER_A,)):
    return {"Authorization": f"Bearer {_token(ent, act=('read', 'write'))}"}


def _raw_dsn():
    return (PG_DSN or "").replace("postgresql+psycopg2://", "postgresql://")


def _add_keyword_column():
    with psycopg2.connect(_raw_dsn()) as c, c.cursor() as cur:
        cur.execute(
            "ALTER TABLE langchain_pg_embedding ADD COLUMN IF NOT EXISTS document_tsv tsvector "
            "GENERATED ALWAYS AS (to_tsvector('english', document)) STORED"
        )
        cur.execute(
            "CREATE INDEX IF NOT EXISTS ix_document_tsv ON langchain_pg_embedding USING gin (document_tsv)"
        )
        c.commit()


class _Recorder:
    """What actually ran for the LAST request (reset() before each call):
    * the dense arm's own result list;
    * the keyword leg's calls and its own result list (appended only when the leg SUCCEEDS; a
      failed keyword leg is swallowed by the route into a dense-only fallback);
    * the Bedrock rerank attempts."""

    def __init__(self, path):
        self.path = path
        self.reset()

    def reset(self):
        self.dense_results = []
        self.keyword_calls = 0
        self.keyword_results = []
        self.rerank_attempts = 0


def _select_path(monkeypatch, store, path):
    """default: the app's defaults + a live keyword leg + Bedrock rerank failing (real fallback).
    dense: hybrid and rerank off. Either way, returns a _Recorder of what ran."""
    from app.routes import document_routes as dr
    from app.services import database as db
    from app.services import reranker
    from app.services.database import PSQLDatabase

    rec = _Recorder(path)
    if path == "dense":
        monkeypatch.setattr(dr, "HYBRID_SEARCH_ENABLED", False)
        monkeypatch.setattr(dr, "RERANK_ENABLED", False)
    else:
        assert (dr.HYBRID_SEARCH_ENABLED, dr.RERANK_ENABLED, reranker.RERANK_ENABLED) == (True, True, True), (
            "precondition: the app defaults are hybrid ON and rerank ON",
            dr.HYBRID_SEARCH_ENABLED, dr.RERANK_ENABLED, reranker.RERANK_ENABLED)
        PSQLDatabase.pool = None
        monkeypatch.setattr(db, "DSN", _raw_dsn(), raising=True)

    real_dense = store.asimilarity_search_with_score_by_vector

    async def _dense(*a, **k):
        out = await real_dense(*a, **k)
        rec.dense_results.append(list(out))
        return out

    monkeypatch.setattr(store, "asimilarity_search_with_score_by_vector", _dense)

    real_kw = dr.keyword_search

    async def _kw(*a, **k):
        # TestClient runs each request on its own loop: rebuild the asyncpg pool per call
        # (same pattern as test_both_legs_expected_passage_and_version / F-ENTITLEMENT-FUSED).
        rec.keyword_calls += 1
        try:
            out = await real_kw(*a, **k)
            rec.keyword_results.append(list(out))
            return out
        finally:
            await PSQLDatabase.close_pool()

    monkeypatch.setattr(dr, "keyword_search", _kw)

    def _bedrock_unreachable(*a, **k):
        rec.rerank_attempts += 1
        raise RuntimeError("SYNTHETIC: Bedrock rerank unreachable in tests")

    monkeypatch.setattr(reranker, "_rerank_sync", _bedrock_unreachable)
    return rec


@pytest.fixture()
def pool_cleanup():
    yield
    from app.services.database import PSQLDatabase

    if PSQLDatabase.pool is not None:
        PSQLDatabase.pool.terminate()
        PSQLDatabase.pool = None


def _embed(client, file_id, name, mime, content, entity, ent=None):
    r = client.post("/embed", data={"file_id": file_id, "entity_id": entity},
                    headers=_write_hdr(ent or (entity,)),
                    files={"file": (name, io.BytesIO(content), mime)})
    assert r.status_code == 200, r.text
    return r


def _store_with(monkeypatch, collection, fmts, path="default"):
    """A fresh real pgvector table holding the named SYNTHETIC fixtures, embedded through the
    REAL /embed route as userA, and a client wired to the chosen path. Returns (store, client,
    recorder)."""
    store = _real_store(monkeypatch, collection)
    _add_keyword_column()
    client = _client(monkeypatch, store)
    monkeypatch.setattr("app.config.vector_store", store, raising=False)
    for fmt in fmts:
        fid, name, mime, make = FIXTURES[fmt]
        _embed(client, fid, name, mime, make(), USER_A)
    rec = _select_path(monkeypatch, store, path)
    return store, client, rec


def _file_leg(client, file_id, query, hdr=None, **extra):
    body = {"file_id": file_id, "query": query, "k": CORE_K, **extra}
    return client.post("/query", json=body, headers=hdr if hdr is not None else _core_hdr())


def _knowledge_leg(client, entity, query, hdr=None, k=CORE_K, args=None):
    body = {"query": query, "k": k, "args": {"cohere": False} if args is None else args}
    return client.post(f"/query/{entity}", json=body, headers=hdr if hdr is not None else _core_hdr())


def _assert_mode(rec, leg, token=None):
    """Which retrieval mode ACTUALLY ran for the last request (see the module docstring)."""
    if rec.path == "dense":
        assert rec.keyword_calls == 0, ("the dense path called the keyword leg", rec.keyword_calls)
        assert rec.rerank_attempts == 0, ("the dense path attempted a rerank", rec.rerank_attempts)
        assert len(rec.dense_results) == 1, ("the dense arm did not run exactly once", len(rec.dense_results))
        return
    assert rec.keyword_calls == 1, (leg, "keyword leg calls", rec.keyword_calls)
    assert len(rec.keyword_results) == 1, (leg, "the keyword leg FAILED (swallowed into a dense-only fallback)")
    rows = rec.keyword_results[0]
    if token is None:
        assert rows, (leg, "the keyword leg matched nothing")
    else:
        assert any(token in (d.page_content or "") for d, _ in rows), (leg, "the keyword leg missed", token, rows)
    expected = 1 if leg == "file" else 0
    assert rec.rerank_attempts == expected, (leg, "Bedrock rerank attempts", rec.rerank_attempts, "expected", expected)


def _legs(client, rec, file_id, query, token):
    """Both routes Core uses for this content, as (name, hits), each checked for the mode that ran."""
    out = []
    for leg, name, call in (
        ("file", "file_leg /query", lambda: _file_leg(client, file_id, query)),
        ("knowledge", "knowledge_leg /query/{entity}", lambda: _knowledge_leg(client, USER_A, query)),
    ):
        rec.reset()
        resp = call()
        assert resp.status_code == 200, (name, resp.text)
        _assert_mode(rec, leg, token)
        out.append((name, resp.json()))
    return out


def _carrying(hits, token):
    return [h for h in hits if token in (h[0].get("page_content") or "")]


def _only(hits, token, where):
    found = _carrying(hits, token)
    assert len(found) == 1, f"{where}: expected exactly one chunk carrying {token!r}, got {len(found)}: {hits!r}"
    return found[0][0]["metadata"]


def _assert_no_ties(rec, hits):
    """The premise the exact-repeat tests stand on. No two candidates may share a dense distance
    or a keyword score, because the store orders by those alone (FILES finding). Fused scores are
    NOT checked."""
    assert rec.dense_results, "precondition: the dense arm's own result was recorded"
    for lst in rec.dense_results:
        d = [s for _, s in lst]
        assert len(set(d)) == len(d), ("tied dense distances in the fixture", d)
    for lst in rec.keyword_results:
        s = [x for _, x in lst]
        assert len(set(s)) == len(s), ("tied keyword scores in the fixture", s)
    if rec.path == "dense":
        w = [h[1] for h in hits]
        assert len(set(w)) == len(w), ("tied wire distances", w)


def _assert_answers(hits, leg, where):
    """The repeated result IS the answer, and it carries its locator."""
    assert _only(hits, DOCX_TOKEN, where).get("block_index") == 1, where
    if leg == "knowledge":
        assert _only(hits, XLSX_TOKEN, where).get("cell_range") == "Revenue!A1:C3", where
        assert _only(hits, MD_TOKEN, where).get("section_index") == 1, where


# =========================================================================== (A) ACCESS
@needs_pg
@pytest.mark.parametrize("path", ["default", "dense"])
def test_A_docx_block_index_reaches_both_core_routes_and_header_footer_carry_none_SYNTHETIC(
        monkeypatch, pool_cleanup, path):
    _, client, rec = _store_with(monkeypatch, "c2p1_access_docx", ["docx"], path)
    q = f"What is the {DOCX_TOKEN} quarterly revenue figure?"
    for name, hits in _legs(client, rec, F_DOCX, q, DOCX_TOKEN):
        # No fabrication on the READ path, checked FIRST so that a stamped aux unit reddens HERE:
        # header/footer units come back WITHOUT the key.
        for tok in (HDR_TOKEN, FTR_TOKEN):
            aux = _only(hits, tok, f"{name} aux {tok}")
            assert "block_index" not in aux, (name, tok, aux)
            assert aux.get("file_id") == F_DOCX, aux
        md = _only(hits, DOCX_TOKEN, name)
        assert md.get("block_index") == 1, (name, md)
        assert (md.get("file_id"), md.get("user_id"), md.get("filename")) == (F_DOCX, USER_A, "report.docx"), md
        body = [h for h in hits
                if not any(t in (h[0].get("page_content") or "") for t in (HDR_TOKEN, FTR_TOKEN))]
        body_idx = {h[0]["metadata"].get("block_index") for h in body}
        assert len(body) == 3 and body_idx == {0, 1, 2}, (name, body_idx, len(body))


@needs_pg
@pytest.mark.parametrize("path", ["default", "dense"])
def test_A_xlsx_cell_range_and_header_reach_both_core_routes_SYNTHETIC(monkeypatch, pool_cleanup, path):
    from openpyxl import load_workbook

    # Fixture precondition: two non-empty sheets and one blank sheet really are in the workbook.
    assert load_workbook(io.BytesIO(xlsx_with_blank_sheet_SYNTHETIC())).sheetnames == ["Revenue", "Costs", "Blank"]
    _, client, rec = _store_with(monkeypatch, "c2p1_access_xlsx", ["xlsx"], path)
    for name, hits in _legs(client, rec, F_XLSX, f"Which region carries {XLSX_TOKEN}?", XLSX_TOKEN):
        md = _only(hits, XLSX_TOKEN, name)
        assert md.get("cell_range") == "Revenue!A1:C3", (name, md)
        assert md.get("header") == ["Region", "Metric", "Value"], (name, md)
        assert md.get("header_row") == 1, (name, md)
        assert md.get("page_name") == "Revenue", (name, md)
        assert (md.get("file_id"), md.get("user_id")) == (F_XLSX, USER_A), md
        # The empty sheet has no occupied extent, emits no chunk, and so is NOT reachable
        # through the query path. Its no-fabrication is pinned at the loader level, in
        # test_xlsx_cell_locator::test_SYNTHETIC_empty_sheet_has_no_cell_range_at_the_locator_level.
        # What IS reachable: no chunk names the blank sheet.
        assert not [h for h in hits if h[0]["metadata"].get("page_name") == "Blank"], (name, hits)
        # Every range present is qualified to the chunk's OWN sheet. This is load-bearing because
        # the fixture carries a SECOND non-empty sheet (Costs).
        qualified = 0
        for h in hits:
            m = h[0]["metadata"]
            if "cell_range" in m:
                assert m["cell_range"].startswith(f"{m.get('page_name')}!"), (name, m)
                qualified += 1
        assert qualified >= 2, (name, "precondition: both non-empty sheets carry a cell_range", qualified)


@needs_pg
@pytest.mark.parametrize("path", ["default", "dense"])
def test_A_md_section_index_and_heading_path_reach_both_core_routes_preamble_carries_none_SYNTHETIC(
        monkeypatch, pool_cleanup, path):
    _, client, rec = _store_with(monkeypatch, "c2p1_access_md", ["md"], path)
    q = f"What does the {MD_TOKEN} revenue section say?"
    for name, hits in _legs(client, rec, F_MD, q, MD_TOKEN):
        md = _only(hits, MD_TOKEN, name)
        assert md.get("section_index") == 1, (name, md)
        # heading_path is a PROPOSED, display-only field pending Core (document_loader.py,
        # MD_HEADING_PATH_KEY); the ADDRESS is section_index. Pinned as emitted, not as contract.
        assert md.get("heading_path") == "Overview > Revenue", (name, md)
        assert (md.get("file_id"), md.get("user_id"), md.get("filename")) == (F_MD, USER_A, "notes.md"), md
        pre = _only(hits, PRE_TOKEN, f"{name} preamble")
        assert "section_index" not in pre and "heading_path" not in pre, (name, pre)


@needs_pg
def test_A_locators_survive_on_keyword_sourced_results_SYNTHETIC(monkeypatch, pool_cleanup):
    """The dense arm is stubbed to contribute nothing (pattern: test_both_legs, keyword-only), so
    every wire Document is built by the keyword leg from its own row. Its locators must survive."""
    store, client, rec = _store_with(monkeypatch, "c2p1_access_kw", ["docx", "xlsx", "md"], "default")
    dense_stub_calls = []

    async def _no_dense(*a, **k):
        dense_stub_calls.append(1)
        return []

    monkeypatch.setattr(store, "asimilarity_search_with_score_by_vector", _no_dense)
    cases = (
        (F_DOCX, DOCX_TOKEN, f"What is the {DOCX_TOKEN} quarterly revenue figure?",
         {"block_index": 1}),
        (F_XLSX, XLSX_TOKEN, f"Which region carries {XLSX_TOKEN}?",
         {"cell_range": "Revenue!A1:C3", "header": ["Region", "Metric", "Value"], "header_row": 1,
          "page_name": "Revenue"}),
        (F_MD, MD_TOKEN, f"What does the {MD_TOKEN} revenue section say?",
         {"section_index": 1, "heading_path": "Overview > Revenue"}),
    )
    for fid, token, q, want in cases:
        for name, hits in _legs(client, rec, fid, q, token):
            md = _only(hits, token, f"keyword-sourced {name}")
            for key, value in want.items():
                assert md.get(key) == value, (name, key, md)
            assert (md.get("file_id"), md.get("user_id")) == (fid, USER_A), md
    assert len(dense_stub_calls) == 2 * len(cases), ("the dense arm must be stubbed on every request",
                                                     len(dense_stub_calls))


@needs_pg
def test_A_core_agent_file_leg_and_two_entity_token_reach_the_locator_SYNTHETIC(monkeypatch, pool_cleanup):
    """Core's agent-attached file sends body entity_id = the agent, with token
    ent = [user, agent] (retrieval.js @091d9729). The same two-entity token on a plain file leg
    exercises the 2-element user filter ($in / = ANY), which a single-entity token never does."""
    _, client, rec = _store_with(monkeypatch, "c2p1_access_agent", ["docx"], "default")
    _fid, name, mime, make = FIXTURES["docx"]
    _embed(client, F_AGENT_DOCX, name, mime, make(), AGENT, ent=(USER_A, AGENT))
    hdr = _core_hdr((USER_A, AGENT))
    q = f"What is the {DOCX_TOKEN} quarterly revenue figure?"

    rec.reset()
    r = _file_leg(client, F_AGENT_DOCX, q, hdr=hdr, entity_id=AGENT)
    assert r.status_code == 200, r.text
    _assert_mode(rec, "file", DOCX_TOKEN)
    hits = r.json()
    md = _only(hits, DOCX_TOKEN, "agent file leg")
    assert (md.get("block_index"), md.get("user_id"), md.get("file_id")) == (1, AGENT, F_AGENT_DOCX), md
    assert {h[0]["metadata"].get("user_id") for h in hits} == {AGENT}, hits

    for file_id, owner in ((F_DOCX, USER_A), (F_AGENT_DOCX, AGENT)):
        rec.reset()
        r = _file_leg(client, file_id, q, hdr=hdr)
        assert r.status_code == 200, r.text
        _assert_mode(rec, "file", DOCX_TOKEN)
        hits = r.json()
        md = _only(hits, DOCX_TOKEN, f"two-entity token, file {file_id}")
        assert (md.get("block_index"), md.get("user_id"), md.get("file_id")) == (1, owner, file_id), md
        assert {h[0]["metadata"].get("file_id") for h in hits} == {file_id}, hits


@needs_pg
def test_A_knowledge_leg_with_rerank_on_keeps_locators_and_caps_at_rerank_top_n_SYNTHETIC(
        monkeypatch, pool_cleanup):
    """Core sends knowledge.args ?? {cohere: false}. For a base whose args lack `cohere`, the
    rerank is ON for its knowledge leg (document_routes._cohere_rerank_enabled). In this image the
    Bedrock call fails, so the real rerank() fallback returns the first RERANK_TOP_N candidates.
    Deterministic: the three answer chunks match the keyword leg and sit in both RRF lists, so they
    outrank every dense-only chunk."""
    from app.routes import document_routes as dr

    _, client, rec = _store_with(monkeypatch, "c2p1_rerank_kn", ["docx", "xlsx", "md"], "default")
    rec.reset()
    r = _knowledge_leg(client, USER_A, Q_REPEAT, args={})
    assert r.status_code == 200, r.text
    hits = r.json()
    assert (rec.keyword_calls, len(rec.keyword_results), rec.rerank_attempts) == (1, 1, 1), (
        rec.keyword_calls, len(rec.keyword_results), rec.rerank_attempts)
    _assert_no_ties(rec, hits)
    scope = len(rec.dense_results[0])
    assert scope > dr.RERANK_TOP_N, ("precondition: more candidates than the rerank cap", scope, dr.RERANK_TOP_N)
    assert len(hits) == dr.RERANK_TOP_N, (len(hits), dr.RERANK_TOP_N)
    _assert_answers(hits, "knowledge", "rerank-on knowledge leg")
    assert _knowledge_leg(client, USER_A, Q_REPEAT, args={}).content == r.content, "the rerank path repeat diverged"


# =========================================================================== (B) ISOLATION
@needs_pg
def test_B_a_file_scoped_query_never_returns_another_file_of_the_same_user_SYNTHETIC(
        monkeypatch, pool_cleanup):
    _, client, _rec = _store_with(monkeypatch, "c2p1_iso_file", ["docx", "md"])
    q = f"What does the {MD_TOKEN} revenue section say?"
    # Positive control: the md file DOES answer this question, so the negative is not vacuous.
    own = _file_leg(client, F_MD, q)
    assert own.status_code == 200 and _carrying(own.json(), MD_TOKEN), own.text
    other = _file_leg(client, F_DOCX, q)
    assert other.status_code == 200, other.text
    hits = other.json()
    assert hits, "precondition: the docx file has chunks of its own to return"
    assert {h[0]["metadata"].get("file_id") for h in hits} == {F_DOCX}, hits
    assert not _carrying(hits, MD_TOKEN), hits
    assert {h[0]["metadata"].get("user_id") for h in hits} == {USER_A}, hits


@needs_pg
def test_B_another_entity_retrieves_nothing_of_user_a_on_any_core_route_SYNTHETIC(
        monkeypatch, pool_cleanup):
    _, client, _rec = _store_with(monkeypatch, "c2p1_iso_entity", ["docx"])
    _fid, name, mime, make = FIXTURES["md"]
    _embed(client, F_B_MD, name, mime, make(), USER_B)       # B owns a file of its own
    q = f"What is the {DOCX_TOKEN} quarterly revenue figure?"
    hdr_b = _core_hdr((USER_B,))

    # Positive controls, SAME routes and query: A's valid token reaches A's answer on both legs.
    assert _carrying(_file_leg(client, F_DOCX, q).json(), DOCX_TOKEN), "file-leg positive control failed"
    r = _knowledge_leg(client, USER_A, q)
    assert r.status_code == 200 and _carrying(r.json(), DOCX_TOKEN), r.text

    r = _file_leg(client, F_DOCX, q, hdr=hdr_b)            # file leg, A's file id, B's token
    assert r.status_code == 200, r.text
    assert r.json() == [], r.json()

    r = _knowledge_leg(client, USER_A, q, hdr=hdr_b)       # knowledge leg naming A
    assert r.status_code == 403, r.text
    assert not isinstance(r.json(), list), r.json()

    r = _knowledge_leg(client, USER_B, q, hdr=hdr_b)       # B's own knowledge leg: B's rows only
    assert r.status_code == 200, r.text
    hits = r.json()
    assert hits, "positive control: B's own knowledge leg returns B's rows"
    assert {h[0]["metadata"].get("user_id") for h in hits} == {USER_B}, hits
    assert {h[0]["metadata"].get("file_id") for h in hits} == {F_B_MD}, hits
    assert not _carrying(hits, DOCX_TOKEN), hits

    r = _file_leg(client, F_DOCX, q, hdr=hdr_b, entity_id=USER_A)   # agent-file shape naming A
    assert r.status_code == 403, r.text
    assert not isinstance(r.json(), list), r.json()


@needs_pg
def test_B_missing_or_invalid_auth_never_falls_back_to_data_SYNTHETIC(monkeypatch, pool_cleanup):
    _, client, _rec = _store_with(monkeypatch, "c2p1_iso_auth", ["docx"])
    q = f"What is the {DOCX_TOKEN} quarterly revenue figure?"
    # Positive controls on BOTH routes the negatives hit: a valid token gets the answer.
    assert _carrying(_file_leg(client, F_DOCX, q).json(), DOCX_TOKEN), "file-leg positive control failed"
    r = _knowledge_leg(client, USER_A, q)
    assert r.status_code == 200 and _carrying(r.json(), DOCX_TOKEN), r.text
    bad = {
        "missing": {"Content-Type": "application/json"},
        "not-bearer": {"Authorization": _token(), "Content-Type": "application/json"},
        "garbage": {"Authorization": "Bearer not.a.jwt", "Content-Type": "application/json"},
        "wrong-secret": {"Authorization": f"Bearer {_token(secret='SYNTHETIC-forged-secret')}",
                         "Content-Type": "application/json"},
        "expired": {"Authorization": f"Bearer {_token(hours=-1)}", "Content-Type": "application/json"},
    }
    for label, hdr in bad.items():
        for route, resp in (("/query", _file_leg(client, F_DOCX, q, hdr=hdr)),
                            ("/query/{entity}", _knowledge_leg(client, USER_A, q, hdr=hdr))):
            assert resp.status_code == 401, (label, route, resp.status_code, resp.text)
            assert not isinstance(resp.json(), list), (label, route, resp.text)
            assert DOCX_TOKEN not in resp.text, (label, route, resp.text)


# =========================================================================== (C) EXACT REPEAT
def _fresh_store(collection):
    """A NEW store object -- new SQLAlchemy engine, new connections -- bound to the SAME
    already-seeded table (no drop, no re-embed)."""
    from langchain_community.vectorstores.pgvector import _get_embedding_collection_store

    from app.services.vector_store.factory import get_vector_store

    store = get_vector_store(_sqlalchemy_dsn(), _DetEmb(), collection, mode="async")
    if store.create_extension:
        store.create_vector_extension()
    store.EmbeddingStore, store.CollectionStore = _get_embedding_collection_store(
        store._embedding_length, use_jsonb=store.use_jsonb)
    store.create_tables_if_not_exists()
    store.create_collection()
    return store


@needs_pg
@pytest.mark.parametrize("path", ["default", "dense"])
def test_C_same_question_twice_ten_times_and_after_a_fresh_store_engine_and_client_is_byte_identical_SYNTHETIC(
        monkeypatch, pool_cleanup, path):
    """A new store, engine and TestClient on the SAME FastAPI app in the SAME process: this does
    not see per-process state (module caches, hash seeds, app.state)."""
    from app.routes import document_routes as dr

    store, client, rec = _store_with(monkeypatch, "c2p1_repeat", ["docx", "xlsx", "md"], path)
    calls = {
        "file": lambda c: _file_leg(c, F_DOCX, Q_REPEAT),
        "knowledge": lambda c: _knowledge_leg(c, USER_A, Q_REPEAT),
    }
    baseline = {}
    for leg, call in calls.items():
        rec.reset()
        r0 = call(client)
        assert r0.status_code == 200, (leg, r0.text)
        hits = r0.json()
        assert len(hits) >= 3, f"{leg}: precondition -- a multi-chunk ordered result, got {len(hits)}"
        assert all(isinstance(h[1], float) for h in hits), hits
        _assert_mode(rec, leg)
        _assert_no_ties(rec, hits)
        _assert_answers(hits, leg, f"{leg} baseline")
        baseline[leg] = r0.content
        r1 = call(client)                                   # asked twice
        assert r1.content == r0.content, f"{leg}: the second identical request diverged"
        for i in range(10):                                 # asked ten more times
            ri = call(client)
            assert ri.status_code == 200, (leg, f"repeat {i + 3}", ri.text)
            assert ri.content == r0.content, f"{leg}: repeat {i + 3} diverged"

    # A new store object, new engine and connections, new TestClient; the old engine is
    # disposed. Same table, nothing re-embedded.
    store._bind.dispose()
    monkeypatch.setattr(dr, "vector_store", _fresh_store("c2p1_repeat"))
    fresh = TestClient(client.app)
    for leg, call in calls.items():
        rf = call(fresh)
        assert rf.status_code == 200, (leg, rf.text)
        assert rf.content == baseline[leg], f"{leg}: a fresh store, engine and client diverged from the baseline"


@needs_pg
@pytest.mark.parametrize("path", ["default", "dense"])
def test_C_a_k_smaller_than_the_scope_cuts_to_the_same_result_every_time_SYNTHETIC(
        monkeypatch, pool_cleanup, path):
    """In production k is smaller than the corpus. With no ties, the cut must be deterministic.
    Where the pipeline makes the cut a plain head of the complete ranking, the cut must equal that
    head: the dense path, and the default file leg (whose fused list is identical at both k).
    The default knowledge leg fuses per-k lists, so only its determinism is asserted."""
    small_k = 4
    _, client, rec = _store_with(monkeypatch, "c2p1_repeat_cut", ["docx", "xlsx", "md"], path)
    for leg in ("file", "knowledge"):
        def call(k, _leg=leg):
            if _leg == "file":
                return _file_leg(client, F_DOCX, Q_REPEAT, k=k)
            return _knowledge_leg(client, USER_A, Q_REPEAT, k=k)

        rec.reset()
        r0 = call(small_k)
        assert r0.status_code == 200, (leg, r0.text)
        hits = r0.json()
        _assert_mode(rec, leg)
        _assert_no_ties(rec, hits)
        full = call(CORE_K).json()
        assert len(full) > small_k, (leg, "precondition: the scope is larger than k", len(full))
        assert len(hits) == small_k, (leg, "the k-cut did not happen", len(hits))
        if path == "dense" or leg == "file":
            assert hits == full[:small_k], (leg, "the k-cut is not the head of the complete ranking")
        for i in range(11):
            ri = call(small_k)
            assert ri.status_code == 200, (leg, f"k={small_k} repeat {i + 2}", ri.text)
            assert ri.content == r0.content, f"{leg} k={small_k}: repeat {i + 2} diverged"
