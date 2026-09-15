"""Mandatory XLSX capability suite for the /embed pipeline (KI-02 SP-01.5).

Excel was the one common Office format that could NEVER be ingested in
production: `UnstructuredExcelLoader` imports `msoffcrypto` at load time and
that package was in no requirements file, so every .xlsx and .xls raised
`ModuleNotFoundError` inside `.load()` and the route answered
`400 "Error during file processing: No module named 'msoffcrypto'"` with zero
vector rows. Reproduced end-to-end on a production-equivalent image (the repo's
own Dockerfile.lite dependency layers) before the fix.

Nothing in this file may skip. The previous XLSX tests were gated on
`msoffcrypto` being importable, so the exact defect that broke production made
its own tests disappear from CI instead of failing it. `test_msoffcrypto_*`
below fails loudly if the dependency is ever dropped again, and every other test
here runs unconditionally.

What is proven:
  * the dependency is pinned in BOTH requirements files (production, not just CI),
  * multi-sheet extraction with exact per-sheet citations,
  * a formula's CACHED value is extracted, and an UNCACHED formula is reported
    explicitly instead of the value going silently missing — never a computed
    substitute,
  * password-protected gets its own `encrypted` verdict, distinct from `corrupt`,
  * a damaged container gets `corrupt`, and an empty workbook stays the existing
    `empty` 422 — all three write zero vector rows,
  * a retried upload re-derives identical chunks (deterministic extraction — NOT
    store-level de-duplication, which needs a real database and belongs to
    SP-01.3), and another tenant's entitlement cannot embed into this entity,
  * the diagnostic formula scan is BOUNDED, and /local/embed reaches the same
    verdicts as /embed — the verdict must not depend on which route was used.

Fixtures are SYNTHETIC and built at test time (openpyxl, plus msoffcrypto's own
encryptor for the password-protected case). No client content. The vector store
is SIMULATED — `AsyncPgVector.aadd_documents` is recorded — so "no rows written"
is asserted against the actual insert attempts.
"""

import datetime
import io
import os
import shutil
import zipfile
from concurrent.futures import ThreadPoolExecutor

import jwt
import pytest
from fastapi.testclient import TestClient

from main import app
from app.services.vector_store.async_pg_vector import AsyncPgVector
from app.utils.document_loader import (
    CorruptDocumentError,
    EncryptedDocumentError,
    SheetExcelLoader,
    get_loader,
)

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_REPO_ROOT = os.path.dirname(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
)
_SECRET = "testsecret"


# ===========================================================================
# Synthetic fixtures
# ===========================================================================


def make_multisheet_workbook(path, cached_total=False):
    """Two sheets; 'Revenue' ends in a =SUM() total row.

    openpyxl writes the formula with NO cached result (it does not calculate),
    which is exactly the uncached case. `cached_total=True` then injects the
    `<v>` element Excel itself would have stored, giving the cached case — the
    only difference between the two fixtures.
    """
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Revenue"
    ws["A1"] = "Region"
    ws["B1"] = "Amount"
    ws["A2"] = "North"
    ws["B2"] = 100
    ws["A3"] = "South"
    ws["B3"] = 200
    ws["A4"] = "Total"
    ws["B4"] = "=SUM(B2:B3)"
    notes = wb.create_sheet("Notes")
    notes["A1"] = "Escalation owner is the regional lead"
    wb.save(path)

    if cached_total:
        tmp = str(path) + ".tmp"
        with zipfile.ZipFile(path) as zin, zipfile.ZipFile(
            tmp, "w", zipfile.ZIP_DEFLATED
        ) as zout:
            for item in zin.infolist():
                payload = zin.read(item.filename)
                if item.filename == "xl/worksheets/sheet1.xml":
                    payload = payload.replace(
                        b"<f>SUM(B2:B3)</f>", b"<f>SUM(B2:B3)</f><v>300</v>"
                    )
                zout.writestr(item, payload)
        shutil.move(tmp, str(path))


def make_empty_workbook(path):
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.title = "Blank"
    wb.save(path)


def make_corrupt_workbook(path):
    """ZIP magic, but not an OOXML package."""
    with open(path, "wb") as f:
        f.write(b"PK\x03\x04 this is not a real office package \x00\x01\x02\x03")


def make_damaged_ole_workbook(path):
    """An OLE2 header with no usable compound-file structure behind it."""
    with open(path, "wb") as f:
        f.write(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 128)


def make_password_protected_workbook(path, tmp_path, password="hunter2"):
    """A REAL encrypted OOXML workbook, produced by msoffcrypto's own encryptor."""
    from msoffcrypto.format.ooxml import OOXMLFile

    plain = tmp_path / "plain-source.xlsx"
    make_multisheet_workbook(str(plain))
    with open(plain, "rb") as src, open(path, "wb") as out:
        OOXMLFile(src).encrypt(password, out)


def load_documents(path):
    loader, known_type, ext = get_loader(os.path.basename(str(path)), XLSX_MIME, str(path))
    assert isinstance(loader, SheetExcelLoader)
    assert known_type is True and ext == "xlsx"
    return loader.load()


# ===========================================================================
# The dependency itself — the regression that broke production
# ===========================================================================


def test_msoffcrypto_is_importable():
    """RED before SP-01.5: this import is what every .xlsx died on.

    No skip marker on purpose. If msoffcrypto ever leaves the image again, this
    fails instead of quietly removing the Excel tests from the run.
    """
    import msoffcrypto  # noqa: F401

    assert hasattr(msoffcrypto, "OfficeFile")


@pytest.mark.parametrize("requirements", ["requirements.txt", "requirements.lite.txt"])
def test_msoffcrypto_is_pinned_in_requirements(requirements):
    """The TEST image having the package proves nothing about production.

    Both images install from a requirements file, so the pin has to be in both
    or the deployed service keeps failing on every spreadsheet.
    """
    with open(os.path.join(_REPO_ROOT, requirements), encoding="utf-8") as f:
        lines = [
            line.strip()
            for line in f
            if line.strip() and not line.strip().startswith("#")
        ]
    pins = [line for line in lines if line.lower().startswith("msoffcrypto-tool")]
    assert pins, f"msoffcrypto-tool missing from {requirements}"
    assert all("==" in pin for pin in pins), f"unpinned in {requirements}: {pins}"


# ===========================================================================
# Extraction + exact sheet citations
# ===========================================================================


def test_multisheet_workbook_extracts_with_sheet_citations(tmp_path):
    """Every Document carries its sheet name and index, and cell values land on
    the sheet they came from — the citation the RAG pipeline must surface."""
    path = tmp_path / "book.xlsx"
    make_multisheet_workbook(str(path))

    docs = load_documents(path)

    assert docs, "a two-sheet workbook must extract at least one Document"
    assert {"Revenue", "Notes"} == {d.metadata.get("page_name") for d in docs}
    for d in docs:
        assert isinstance(d.metadata.get("page_number"), int)

    by_sheet = {}
    for d in docs:
        by_sheet.setdefault(d.metadata["page_name"], []).append(d.page_content)
    revenue = " ".join(by_sheet["Revenue"])
    notes = " ".join(by_sheet["Notes"])
    assert "North" in revenue and "100" in revenue
    assert "Escalation owner" in notes
    assert "Escalation owner" not in revenue


def test_cached_formula_value_is_extracted(tmp_path):
    """A formula Excel has already calculated stores its result in the file, and
    that result is real extracted content."""
    path = tmp_path / "cached.xlsx"
    make_multisheet_workbook(str(path), cached_total=True)

    docs = load_documents(path)
    revenue = " ".join(
        d.page_content for d in docs if d.metadata.get("page_name") == "Revenue"
    )
    assert "300" in revenue
    assert all(d.metadata.get("formula_scan") == "complete" for d in docs)
    assert all("formula_uncached" not in d.metadata for d in docs)


def test_uncached_formula_is_reported_and_never_fabricated(tmp_path):
    """The honesty case.

    A formula with no cached result extracts as a label with the number missing
    ("Total" and nothing else). We must NOT compute 300 and present it as
    content the file contains — so the value stays absent and the omission is
    reported explicitly on the affected sheet.
    """
    path = tmp_path / "uncached.xlsx"
    make_multisheet_workbook(str(path), cached_total=False)

    docs = load_documents(path)
    revenue_docs = [d for d in docs if d.metadata.get("page_name") == "Revenue"]
    revenue = " ".join(d.page_content for d in revenue_docs)

    assert "Total" in revenue
    assert "300" not in revenue, "the uncalculated total must never be invented"
    assert all(d.metadata["formula_scan"] == "complete" for d in revenue_docs)
    assert all(d.metadata["formula_uncached"] == 1 for d in revenue_docs)
    assert all(d.metadata["formula_uncached_cells"] == ["B4"] for d in revenue_docs)

    # The sheet with no formulas is not tarred with the same brush.
    for d in docs:
        if d.metadata.get("page_name") == "Notes":
            assert "formula_uncached" not in d.metadata


def test_formula_scan_is_unavailable_for_a_non_ooxml_container(tmp_path):
    """Real branch, no mocking: openpyxl cannot open a legacy OLE2 workbook, so
    the scan reports `unavailable` instead of a zero it never verified."""
    path = tmp_path / "legacy.xls"
    make_damaged_ole_workbook(str(path))

    cells, scan_status = SheetExcelLoader(str(path))._uncached_formulas()

    assert scan_status == "unavailable"
    assert cells == {}


def test_formula_scan_failure_reports_unavailable_not_a_clean_bill(
    tmp_path, monkeypatch
):
    """SIMULATED scan failure: the formula pass is made to raise.

    Only the `data_only=False` call is broken — that is the formula pass this
    scan owns; pandas (inside unstructured) always reads with `data_only=True`,
    so the real extraction still runs. Two things must hold: the workbook still
    ingests (a diagnostic scan must never cost us the content), and the status
    says `unavailable` rather than reporting zero uncached formulas we never
    actually checked for.
    """
    path = tmp_path / "book.xlsx"
    make_multisheet_workbook(str(path), cached_total=False)

    import openpyxl

    real_load_workbook = openpyxl.load_workbook

    def selective_failure(*args, **kwargs):
        if kwargs.get("data_only") is False:
            raise RuntimeError("simulated formula-pass failure")
        return real_load_workbook(*args, **kwargs)

    monkeypatch.setattr(openpyxl, "load_workbook", selective_failure)

    docs = load_documents(path)
    assert docs, "a failed formula scan must not lose the extracted content"
    assert all(d.metadata["formula_scan"] == "unavailable" for d in docs)
    assert all("formula_uncached" not in d.metadata for d in docs)


# ===========================================================================
# Terminal verdicts — encrypted is NOT corrupt
# ===========================================================================


def test_password_protected_workbook_is_encrypted_not_corrupt(tmp_path):
    """A password-protected workbook is intact; the reader just lacks the key.

    Telling the uploader it is "damaged" would send them to fix a file that has
    nothing wrong with it.
    """
    path = tmp_path / "protected.xlsx"
    make_password_protected_workbook(path, tmp_path)

    with pytest.raises(EncryptedDocumentError) as exc:
        load_documents(path)

    assert exc.value.verdict == "encrypted"
    assert not isinstance(exc.value, CorruptDocumentError)
    assert "password" in str(exc.value).lower()


def test_corrupt_workbook_is_corrupt_not_encrypted(tmp_path):
    path = tmp_path / "corrupt.xlsx"
    make_corrupt_workbook(str(path))

    with pytest.raises(CorruptDocumentError) as exc:
        load_documents(path)

    assert exc.value.verdict == "corrupt"
    assert not isinstance(exc.value, EncryptedDocumentError)


def test_damaged_ole_container_is_corrupt_not_encrypted(tmp_path):
    """An OLE2 header is what an encrypted workbook looks like from outside, so
    the verdict must come from actually reading the container, not the magic."""
    path = tmp_path / "damaged.xlsx"
    make_damaged_ole_workbook(str(path))

    with pytest.raises(CorruptDocumentError) as exc:
        load_documents(path)

    assert exc.value.verdict == "corrupt"


def test_empty_workbook_extracts_nothing_rather_than_raising(tmp_path):
    """An empty workbook is not a failure at the parser — it is caught one level
    up by the empty-extraction guard, which is asserted at the route below."""
    path = tmp_path / "empty.xlsx"
    make_empty_workbook(str(path))

    assert load_documents(path) == []


# ===========================================================================
# Route level — status, receipt, and what does NOT get written
# ===========================================================================


def _hdr(ent, act, tid="tenantA", uid="testuser"):
    os.environ["JWT_SECRET"] = _SECRET
    payload = {
        "id": uid,
        "tid": tid,
        "ent": ent,
        "act": act,
        "exp": datetime.datetime.now(datetime.timezone.utc)
        + datetime.timedelta(hours=1),
    }
    return {"Authorization": f"Bearer {jwt.encode(payload, _SECRET, algorithm='HS256')}"}


@pytest.fixture()
def client(monkeypatch):
    """TestClient with the vector store SIMULATED and every insert recorded."""
    os.environ["JWT_SECRET"] = _SECRET
    if not hasattr(app.state, "thread_pool") or app.state.thread_pool is None:
        app.state.thread_pool = ThreadPoolExecutor(max_workers=2)

    added = []

    async def recording_aadd(self, docs, ids=None, executor=None):
        added.append(list(docs))
        return ids

    async def dummy_delete(
        self,
        ids=None,
        collection_only=False,
        user_id=None,
        document_origin_type=None,
        subscription_id=None,
        executor=None,
    ):
        return None

    monkeypatch.setattr(AsyncPgVector, "aadd_documents", recording_aadd)
    monkeypatch.setattr(AsyncPgVector, "delete", dummy_delete)

    test_client = TestClient(app)
    test_client.inserted_batches = added  # type: ignore[attr-defined]
    return test_client


def _embed(client, filename, content, *, file_id="f-xlsx", entity_id="userA", ent=None):
    return client.post(
        "/embed",
        data={"file_id": file_id, "entity_id": entity_id},
        files={"file": (filename, io.BytesIO(content), XLSX_MIME)},
        headers=_hdr(ent=ent if ent is not None else ["userA"], act=["write"]),
    )


def test_embed_multisheet_xlsx_succeeds_with_sheet_receipt(client, tmp_path):
    """The end-to-end case that was impossible before: a real workbook ingests,
    writes rows, and reports sheet-level extraction."""
    path = tmp_path / "book.xlsx"
    make_multisheet_workbook(str(path), cached_total=True)

    r = _embed(client, "book.xlsx", path.read_bytes())

    assert r.status_code == 200, r.text
    body = r.json()
    assert body["status"] is True
    receipt = body["extraction"]
    assert receipt["locator_kind"] == "sheet"
    assert receipt["status"] == "complete"
    assert receipt["units_extracted"] == 2
    assert receipt["formulas"]["scan"] == "complete"
    assert receipt["formulas"]["uncached_cells_total"] == 0
    assert sum(len(b) for b in client.inserted_batches) >= 1


def test_embed_reports_uncached_formulas_on_the_success_path(client, tmp_path):
    """A workbook whose totals were never calculated still ingests — but the
    caller is told which values are missing, on the 200 response."""
    path = tmp_path / "uncached.xlsx"
    make_multisheet_workbook(str(path), cached_total=False)

    r = _embed(client, "uncached.xlsx", path.read_bytes())

    assert r.status_code == 200, r.text
    formulas = r.json()["extraction"]["formulas"]
    assert formulas["scan"] == "complete"
    assert formulas["uncached_cells_total"] == 1
    assert formulas["uncached"] == [{"locator": "Revenue", "cells": ["B4"]}]
    stored = " ".join(
        d.page_content for batch in client.inserted_batches for d in batch
    )
    assert "300" not in stored, "an uncalculated total must never reach the store"


def test_embed_password_protected_xlsx_is_422_encrypted_with_no_rows(
    client, tmp_path
):
    """Honest and useful: the verdict is machine-readable AND the message tells
    the uploader what to do. Nothing is stored."""
    path = tmp_path / "protected.xlsx"
    make_password_protected_workbook(path, tmp_path)

    r = _embed(client, "protected.xlsx", path.read_bytes())

    assert r.status_code == 422, r.text
    detail = r.json()["detail"]
    assert detail["extraction"]["verdict"] == "encrypted"
    assert detail["extraction"]["status"] == "unsupported"
    assert "protected.xlsx" in detail["message"]
    assert "password" in detail["message"].lower()
    assert client.inserted_batches == []


def test_embed_corrupt_xlsx_is_422_corrupt_with_no_rows(client, tmp_path):
    path = tmp_path / "corrupt.xlsx"
    make_corrupt_workbook(str(path))

    r = _embed(client, "corrupt.xlsx", path.read_bytes())

    assert r.status_code == 422, r.text
    assert r.json()["detail"]["extraction"]["verdict"] == "corrupt"
    assert client.inserted_batches == []


def test_embed_empty_xlsx_is_422_empty_with_no_rows(client, tmp_path):
    """An empty workbook keeps the existing empty-extraction verdict; the new
    verdicts must not swallow it into a generic failure."""
    path = tmp_path / "empty.xlsx"
    make_empty_workbook(str(path))

    r = _embed(client, "empty.xlsx", path.read_bytes())

    assert r.status_code == 422, r.text
    detail = r.json()["detail"]
    assert detail["extraction"]["status"] == "empty"
    assert "verdict" not in detail["extraction"]
    assert client.inserted_batches == []


def test_retrying_the_same_upload_re_derives_identical_chunks(client, tmp_path):
    """A retry after a lost response must not make the second answer differ from
    the first, and must not produce a second, DIVERGENT set of chunks.

    Scope, stated plainly: the store here is SIMULATED and does not de-duplicate,
    so this proves the extraction is deterministic, NOT that the real pgvector
    path replaces rows for a repeated file_id. Store-level idempotence belongs to
    SP-01.3's initialize/finalize work and needs a real database.
    """
    path = tmp_path / "book.xlsx"
    make_multisheet_workbook(str(path), cached_total=True)
    content = path.read_bytes()

    first = _embed(client, "book.xlsx", content, file_id="f-retry")
    batches_after_first = [
        [d.page_content for d in batch] for batch in client.inserted_batches
    ]
    second = _embed(client, "book.xlsx", content, file_id="f-retry")
    batches_after_second = [
        [d.page_content for d in batch] for batch in client.inserted_batches
    ]

    assert first.status_code == 200 and second.status_code == 200
    assert first.json()["extraction"] == second.json()["extraction"]
    replay = batches_after_second[len(batches_after_first):]
    assert replay == batches_after_first, "the retry must re-derive the same chunks"


def test_retrying_a_rejected_upload_still_writes_nothing(client, tmp_path):
    """A terminal verdict is terminal: retrying never turns it into a partial
    write."""
    path = tmp_path / "protected.xlsx"
    make_password_protected_workbook(path, tmp_path)
    content = path.read_bytes()

    first = _embed(client, "protected.xlsx", content, file_id="f-enc")
    second = _embed(client, "protected.xlsx", content, file_id="f-enc")

    assert first.status_code == 422 and second.status_code == 422
    assert first.json()["detail"]["extraction"]["verdict"] == "encrypted"
    assert second.json()["detail"]["extraction"]["verdict"] == "encrypted"
    assert client.inserted_batches == []


def test_embed_into_another_entity_is_denied_with_no_rows(client, tmp_path):
    """A valid workbook and a valid token are not enough: the entitlement must
    cover the entity being written to."""
    path = tmp_path / "book.xlsx"
    make_multisheet_workbook(str(path), cached_total=True)

    r = _embed(
        client,
        "book.xlsx",
        path.read_bytes(),
        entity_id="userB",
        ent=["userA"],
    )

    assert r.status_code == 403, r.text
    assert client.inserted_batches == []


# ===========================================================================
# Bounded processing — the diagnostic scan must never be what makes an upload
# expensive (reviewer WPSP1-5-R NOTE-1)
# ===========================================================================


def test_formula_scan_is_skipped_over_the_size_bound(tmp_path, monkeypatch):
    """A workbook past the byte bound is still extracted; only the extra
    diagnostic passes are dropped, and the status says so."""
    path = tmp_path / "book.xlsx"
    make_multisheet_workbook(str(path), cached_total=False)
    monkeypatch.setattr(SheetExcelLoader, "_MAX_SCAN_BYTES", 10)

    docs = load_documents(path)

    assert docs, "the size bound must cost us the scan, never the content"
    assert all(d.metadata["formula_scan"] == "unavailable" for d in docs)
    assert all("formula_uncached" not in d.metadata for d in docs)


def test_formula_scan_is_abandoned_over_the_cell_bound(tmp_path, monkeypatch):
    """Same for a workbook with too many cells to walk: stop and say
    `unavailable`, rather than report the formulas found before giving up as if
    they were the whole answer."""
    path = tmp_path / "book.xlsx"
    make_multisheet_workbook(str(path), cached_total=False)
    monkeypatch.setattr(SheetExcelLoader, "_MAX_SCAN_CELLS", 1)

    cells, scan_status = SheetExcelLoader(str(path))._uncached_formulas()

    assert scan_status == "unavailable"
    assert cells == {}


# ===========================================================================
# /local/embed reaches the same verdicts (reviewer WPSP1-5-R MINOR-1)
# ===========================================================================


def _local_embed(client, filename, content, *, file_id="f-local", entity_id="userA"):
    """Write the bytes where /local/embed expects them and embed by path."""
    from app.config import RAG_UPLOAD_DIR

    target_dir = os.path.join(RAG_UPLOAD_DIR, entity_id)
    os.makedirs(target_dir, exist_ok=True)
    with open(os.path.join(target_dir, filename), "wb") as f:
        f.write(content)
    return client.post(
        "/local/embed",
        json={
            "filepath": f"{entity_id}/{filename}",
            "filename": filename,
            "file_content_type": XLSX_MIME,
            "file_id": file_id,
        },
        params={"entity_id": entity_id},
        headers=_hdr(ent=[entity_id], act=["write"]),
    )


def test_local_embed_multisheet_xlsx_succeeds_with_sheet_receipt(client, tmp_path):
    path = tmp_path / "book.xlsx"
    make_multisheet_workbook(str(path), cached_total=True)

    r = _local_embed(client, "local-book.xlsx", path.read_bytes())

    assert r.status_code == 200, r.text
    receipt = r.json()["extraction"]
    assert receipt["locator_kind"] == "sheet"
    assert receipt["formulas"]["scan"] == "complete"
    assert sum(len(b) for b in client.inserted_batches) >= 1


def test_local_embed_password_protected_xlsx_gets_the_same_verdict(client, tmp_path):
    """The verdict must not depend on which embed route was used. Before the
    shared-seam fix this route answered a generic 400 with no verdict while
    /embed answered 422 `encrypted` for the very same bytes."""
    path = tmp_path / "protected.xlsx"
    make_password_protected_workbook(path, tmp_path)

    r = _local_embed(client, "local-protected.xlsx", path.read_bytes())

    assert r.status_code == 422, r.text
    detail = r.json()["detail"]
    assert detail["extraction"]["verdict"] == "encrypted"
    assert "password" in detail["message"].lower()
    assert client.inserted_batches == []


def test_local_embed_corrupt_xlsx_gets_the_same_verdict(client, tmp_path):
    path = tmp_path / "corrupt.xlsx"
    make_corrupt_workbook(str(path))

    r = _local_embed(client, "local-corrupt.xlsx", path.read_bytes())

    assert r.status_code == 422, r.text
    assert r.json()["detail"]["extraction"]["verdict"] == "corrupt"
    assert client.inserted_batches == []
