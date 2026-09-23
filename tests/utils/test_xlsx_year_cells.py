"""A year held in a workbook cell is extracted as that year (F-XLSX-YEAR-FORMAT).

MEASURED on origin/main a4b47a6 through the real loader (get_loader ->
SheetExcelLoader -> UnstructuredExcelLoader) with a synthetic workbook:

    int 2015            -> "2015"                 faithful
    float 2015.0        -> "2015"                 faithful
    text "2018"         -> "2018"                 faithful
    "Revenue 2019"      -> "Revenue 2019"         faithful
    date, format yyyy   -> "2016-01-01 00:00:00"  WRONG: Excel shows 2016
    date, default fmt   -> "2017-03-04 00:00:00"  a real date; left as is

A date cell formatted to show only the year is what a spreadsheet author gets
when they want a year column typed as a date. The parser rendered it as a full
timestamp, so the indexed text claimed a day and a month (1 January) the source
never shows. The fix parses a copy in which year-only date cells hold the
integer year; every workbook without such cells is parsed from the original.
How other date formats are displayed (the " 00:00:00" suffix) is an OPEN
display choice, not decided here, and is pinned only so a change is deliberate.

F-XLSX-YEAR-FORMAT-DESTROYS-DATE (2026-09-23) -- THE OPEN CHOICE, CLOSED
-----------------------------------------------------------------------
The line above ("an OPEN display choice, not decided here") was left open when this
card completed, and nothing ever closed it. It became a default by attrition, and the
default destroys data: the rewrite gates on the cell's FORMAT, not on the value's
precision, so a genuine 30 June 2016 displayed as `2016` enters the index as `2016`.
A fiscal year end becomes unanswerable and a mid-year reporter is indistinguishable
from a calendar-year one.

RULED: keep the behaviour, disclose it. In the common case the author typed only a
year and this pass removes a January-the-first timestamp the source never showed --
usually right, occasionally wrong, which is the inverse of a heuristic worth removing.
What was missing is that it was the only content-changing step in this loader that
left no trace on success.

The disclosure names WHAT WAS LOST rather than merely that something happened:
`extraction.dates.reduced` lists the cells whose month and day were discarded, and is
ABSENT when nothing was. A flag on every rewritten year cell would be true on
thousands where nothing was lost, and an alarm that fires on the normal case is one
readers learn to ignore -- which would put the rare destructive case back where it
started.

Stated as a fact about the VALUE, never about intent: rag_api cannot know what the
author typed, only whether anything other than the January-the-first default was
present.

Fixtures are SYNTHETIC, built at test time. No client content.
"""

import datetime
import os
import shutil
import zipfile

import pytest

import app.utils.document_loader as document_loader
from app.utils.document_loader import SheetExcelLoader
from tests.utils.test_xlsx_capability import (  # noqa: F401 - pytest fixture
    _embed,
    client,
    load_documents,
)


def make_year_workbook(path, *, year_format="yyyy", with_year_cell=True, formula=None,
                       year_cell_value=None):
    """One sheet of labelled values covering every way a year can be held."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Years"
    ws.append(["Label", "Value"])
    ws.append(["int_year", 2015])
    ws.append(["float_year", 2015.0])
    ws.append(["text_year", "2018"])
    ws.append(["Revenue 2019", 1234.5])
    ws.append(["date_default", datetime.date(2017, 3, 4)])
    if with_year_cell:
        ws.append(["date_yyyy", year_cell_value or datetime.datetime(2016, 1, 1)])
        ws.cell(row=ws.max_row, column=2).number_format = year_format
    if formula:
        ws.append(["Total", formula])
    wb.save(path)


def inject_cached_value(path, formula_xml, value):
    """Store the result Excel would have cached for a formula openpyxl wrote bare."""
    tmp = str(path) + ".tmp"
    with zipfile.ZipFile(str(path)) as zin, zipfile.ZipFile(tmp, "w") as zout:
        for item in zin.infolist():
            payload = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                assert formula_xml in payload
                payload = payload.replace(formula_xml, formula_xml + value)
            zout.writestr(item, payload)
    shutil.move(tmp, str(path))


def sheet_text(path):
    docs = load_documents(path)
    assert len(docs) == 1
    return docs[0]


# ===========================================================================
# The defect
# ===========================================================================


def test_year_only_date_cell_is_extracted_as_the_year(tmp_path):
    path = tmp_path / "years.xlsx"
    make_year_workbook(str(path))

    doc = sheet_text(path)

    assert "date_yyyy 2016" in doc.page_content, doc.page_content
    assert "2016-01-01" not in doc.page_content
    assert "<td>2016</td>" in doc.metadata["text_as_html"]
    assert "2016-01-01" not in doc.metadata["text_as_html"]


@pytest.mark.parametrize("year_format", ["yy", "YYYY", "[$-409]yyyy;@"])
def test_every_year_only_format_is_extracted_as_the_year(tmp_path, year_format):
    path = tmp_path / "years.xlsx"
    make_year_workbook(str(path), year_format=year_format)

    content = sheet_text(path).page_content

    assert "date_yyyy 2016" in content, content
    assert "2016-01-01" not in content


def test_stored_chunk_carries_the_year_through_the_route(client, tmp_path):
    path = tmp_path / "years.xlsx"
    make_year_workbook(str(path))

    r = _embed(client, "years.xlsx", path.read_bytes(), file_id="f-years")

    assert r.status_code == 200, r.text
    stored = " ".join(
        d.page_content for batch in client.inserted_batches for d in batch
    )
    assert "date_yyyy 2016" in stored, stored
    assert "2016-01-01" not in stored
    assert all(
        d.metadata.get("filename") == "years.xlsx"
        for batch in client.inserted_batches
        for d in batch
    )


# ===========================================================================
# What must NOT change
# ===========================================================================


def test_numeric_and_text_years_are_unchanged(tmp_path):
    """Correct on main before the fix; pinned so the fix cannot disturb them."""
    for with_year_cell in (False, True):
        path = tmp_path / f"years-{with_year_cell}.xlsx"
        make_year_workbook(str(path), with_year_cell=with_year_cell)

        content = sheet_text(path).page_content

        assert "int_year 2015 " in content, content
        assert "float_year 2015 " in content
        assert "2015.0" not in content
        assert "text_year 2018 " in content
        assert "Revenue 2019 1234.5" in content


def test_a_full_date_with_a_normal_format_is_never_reduced_to_its_year(tmp_path):
    """A date shown in full keeps its day and month, including in a workbook that also
    has a year-only cell (the copy path).

    RENAMED. This was `test_a_full_date_is_never_reduced_to_its_year`, which claimed a
    property its fixture never built. The rewrite gates on the cell's FORMAT, not on
    the value's precision, so a real date displayed as `2016` IS reduced -- and the
    only date this fixture builds is `date_default`, which carries a normal format.
    Anyone grepping to find out whether the destructive case was guarded found a test
    whose name said yes.

    The name now states what the fixture actually proves. The case the old name implied
    is covered -- with its loss DISCLOSED rather than prevented, which is the ruling --
    by `test_a_real_date_behind_a_year_only_format_is_reduced_and_the_loss_is_disclosed`.
    """
    path = tmp_path / "years.xlsx"
    make_year_workbook(str(path))

    content = sheet_text(path).page_content

    assert "date_default 2017-03-04" in content, content


def test_workbook_without_year_only_cells_is_parsed_from_the_original(
    tmp_path, monkeypatch
):
    parsed = []
    real = document_loader.UnstructuredExcelLoader

    def spy(file_path, *args, **kwargs):
        parsed.append(file_path)
        return real(file_path, *args, **kwargs)

    monkeypatch.setattr(document_loader, "UnstructuredExcelLoader", spy)

    plain = tmp_path / "plain.xlsx"
    make_year_workbook(str(plain), with_year_cell=False)
    load_documents(plain)
    assert parsed == [str(plain)]

    parsed.clear()
    years = tmp_path / "years.xlsx"
    make_year_workbook(str(years))
    docs = load_documents(years)
    assert len(parsed) == 1 and parsed[0] != str(years)
    assert not os.path.exists(parsed[0]), "working copy must be removed"
    # Provenance names the uploaded file, never the working copy.
    assert docs[0].metadata["source"] == str(years)
    assert docs[0].metadata["file_directory"] == str(tmp_path)
    assert docs[0].metadata["filename"] == "years.xlsx"


def test_cached_formula_value_survives_the_year_pass(tmp_path):
    """The copy is written from cached values; a cached total must still appear."""
    path = tmp_path / "years.xlsx"
    make_year_workbook(str(path), formula="=SUM(B2:B3)")
    inject_cached_value(path, b"<f>SUM(B2:B3)</f>", b"<v>4030</v>")

    doc = sheet_text(path)

    assert "Total 4030" in doc.page_content, doc.page_content
    assert "date_yyyy 2016" in doc.page_content
    assert doc.metadata["formula_scan"] == "complete"
    assert "formula_uncached" not in doc.metadata


def test_uncached_formula_is_still_reported_not_fabricated(tmp_path):
    """The uncached-formula scan reads the ORIGINAL file, so the year pass can
    neither hide an uncached formula nor invent its value."""
    path = tmp_path / "years.xlsx"
    make_year_workbook(str(path), formula="=SUM(B2:B3)")

    # A row with an empty value cell is split off into its own element; that
    # split is existing parser behaviour, so read the sheet as a whole.
    docs = load_documents(path)
    content = " ".join(d.page_content for d in docs)

    assert "date_yyyy 2016" in content, content
    assert "4030" not in content
    for doc in docs:
        assert doc.metadata["formula_scan"] == "complete"
        assert doc.metadata["formula_uncached"] == 1
        assert doc.metadata["formula_uncached_cells"] == ["B8"]


def test_year_pass_failure_parses_the_original(tmp_path, monkeypatch):
    """A failing year pass is never fatal: the workbook still ingests."""
    path = tmp_path / "years.xlsx"
    make_year_workbook(str(path))

    import openpyxl

    def boom(*a, **k):
        raise RuntimeError("synthetic openpyxl failure")

    monkeypatch.setattr(openpyxl, "load_workbook", boom)
    loader = SheetExcelLoader(str(path))
    assert loader._year_only_copy(str(tmp_path)) is None


@pytest.mark.parametrize(
    "fmt,expected",
    [
        ("yyyy", True),
        ("yy", True),
        ("YYYY", True),
        ("[$-409]yyyy;@", True),
        ("[Red]yyyy", True),
        ("yyyy-mm-dd", False),
        ("mmm yyyy", False),
        ("mm-dd-yy", False),
        ("0", False),
        ("General", False),
        (None, False),
    ],
)
def test_year_only_format_recognition(fmt, expected):
    assert SheetExcelLoader._is_year_only_format(fmt) is expected


# ===========================================================================
# F-XLSX-YEAR-FORMAT-DESTROYS-DATE — the case this file's own name promised
# ===========================================================================


def test_a_real_date_behind_a_year_only_format_is_reduced_and_the_loss_is_disclosed(
    tmp_path,
):
    """RED-FIRST for F-XLSX-YEAR-FORMAT-DESTROYS-DATE.

    The rewrite gates on the FORMAT, not on the value's precision, so a genuine
    30 June 2016 displayed as `2016` enters the index as `2016` and its month and day
    are gone. "What is this entity's fiscal year end" becomes unanswerable, and a
    mid-year fiscal reporter is indistinguishable from a calendar-year one.

    The behaviour is KEPT -- in the common case the author typed only a year and this
    pass removes a January-the-first timestamp the source never showed. What was
    missing is that it is the only content-changing step in this module that left no
    trace on success. So the requirement is disclosure, not removal.

    The disclosure must distinguish the two cases. A flag that fires on every year cell
    is an alarm on the normal case: it would be true on thousands of cells where
    nothing was lost, readers would learn to ignore it, and the rare destructive case
    would be back where it started. The distinguishing fact is available at the rewrite
    site, because the value still has its month and day at that moment.
    """
    path = tmp_path / "fiscal.xlsx"
    make_year_workbook(
        str(path), year_cell_value=datetime.datetime(2016, 6, 30)
    )

    doc = sheet_text(path)

    # The reduction itself is intended and stays.
    assert "date_yyyy 2016" in doc.page_content, doc.page_content

    # ...but it must no longer be silent.
    assert doc.metadata.get("date_display_reduced"), (
        "a real date (2016-06-30) was reduced to its year with NO disclosure. The "
        "receipt cannot tell a reader that month and day were discarded, and a year "
        "that was rewritten is indistinguishable from a year that was always a year. "
        "metadata=%r" % (sorted(doc.metadata),))


def test_a_january_first_year_cell_is_not_reported_as_a_loss(tmp_path):
    """THE OTHER HALF, and the reason the disclosure is not a blanket flag.

    A cell holding 2016-01-01 behind a year-only format carries no information beyond
    the year, so rewriting it discards nothing. Reporting it as a loss would fire the
    alarm on the normal case and teach readers to ignore it.

    Stated as a fact about the VALUE, not about intent: rag_api cannot know what the
    author typed, only whether month and day were anything other than the default.
    """
    path = tmp_path / "calendar.xlsx"
    make_year_workbook(str(path))  # default year cell is 2016-01-01

    doc = sheet_text(path)

    assert "date_yyyy 2016" in doc.page_content, doc.page_content
    assert not doc.metadata.get("date_display_reduced"), (
        "a 2016-01-01 cell was reported as a lossy reduction; nothing was discarded, "
        "and an alarm that fires on the normal case is worse than no alarm. "
        "metadata=%r" % (doc.metadata,))


def test_the_receipt_names_the_cells_whose_date_was_reduced(client, tmp_path):
    """THE DISCLOSURE ON THE WIRE, not just in loader metadata.

    A stamp a consumer never receives is not a disclosure. The receipt is what /embed
    returns, so this is the surface that decides whether anyone can learn the month and
    day were dropped.
    """
    path = tmp_path / "fiscal.xlsx"
    make_year_workbook(str(path), year_cell_value=datetime.datetime(2016, 6, 30))

    r = _embed(client, "fiscal.xlsx", path.read_bytes(), file_id="f-fiscal")
    assert r.status_code == 200, r.text

    dates = r.json()["extraction"].get("dates")
    assert dates, (
        "the receipt carries no `dates` block, so nothing on the wire says a real date "
        "was reduced to its year. extraction=%r" % (r.json()["extraction"],))
    assert dates["display_normalised"] == 1
    assert dates["reduced"] == ["Years!B7"], dates


def test_a_lossless_year_pass_reports_no_reduced_cells(client, tmp_path):
    """The block appears (the pass ran) but names nothing, because nothing was lost.

    `reduced` is ABSENT rather than an empty list. An empty list reads as a measurement
    that found none, which is a different claim from "every rewrite was lossless", and
    it is the shape that trains a reader to skim past the key.
    """
    path = tmp_path / "calendar.xlsx"
    make_year_workbook(str(path))  # default year cell is 2016-01-01

    r = _embed(client, "calendar.xlsx", path.read_bytes(), file_id="f-cal")
    assert r.status_code == 200, r.text

    dates = r.json()["extraction"].get("dates")
    assert dates == {"display_normalised": 1}, dates


def test_a_workbook_with_no_year_cells_has_no_dates_block_at_all(client, tmp_path):
    """Receipt shape is unchanged for every file the pass does not touch.

    An additive field that appears on receipts it has nothing to say about is not
    additive in practice: every consumer has to learn to ignore it.
    """
    path = tmp_path / "plain.xlsx"
    make_year_workbook(str(path), with_year_cell=False)

    r = _embed(client, "plain.xlsx", path.read_bytes(), file_id="f-plain")
    assert r.status_code == 200, r.text
    assert "dates" not in r.json()["extraction"], r.json()["extraction"]
