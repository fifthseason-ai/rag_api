"""XLSX finer-than-sheet locators + merged-cell resolution + header row (card E3).

MEASURED base: origin/main 5816e133. Today SheetExcelLoader emits one Document per
sheet carrying `page_name` (sheet) as its only per-unit locator. Card E3 adds:

  * a finer-than-sheet locator -- a sheet-qualified CELL-RANGE (e.g. "Revenue!A1:C5")
    naming the occupied extent -- carried ON TOP of `page_name` on every sheet chunk,
  * MERGED-CELL resolution: a merged range's anchor value is filled across the range,
    so the value is not lost and no phantom empty cells are produced,
  * the HEADER ROW detected and carried as a unit field (`header`, `header_row`).

WHY CELL-RANGE, NOT ROW-LEVEL: the loader wraps UnstructuredExcelLoader(mode="elements"),
which emits ONE element per sheet (pinned by test_xlsx_year_cells.sheet_text's
`len(docs) == 1`). Emitting one Document per ROW would change XLSX chunk granularity
corpus-wide -- a far larger consumer-visible change than the locator itself -- and lose
the per-sheet `text_as_html` the year-cell suite asserts. A sheet-qualified cell-range is
additive metadata on the existing per-sheet chunk: finer than the unbounded sheet
reference, one range per chunk.

PENDING PLACEMENT: the `_UNIT_LOCATOR_KEYS` registration is NOT in this branch (it is the
FILES-lead's placement ruling; see the E3 ASK). Until it lands, an XLSX chunk carries BOTH
`page_name` and `cell_range` but the receipt still advertises locator_kind `sheet`
(Option 1 semantics). These tests read the EXPECTED locator_kind from the module so that
when the placement lands, only that value changes -- the tests do not hardcode "sheet".

Fixtures are SYNTHETIC, built at test time with openpyxl. No client content.

RED-FIRST: the deterministic unit tests below fail if the E3 loader code is reverted --
see the mutation table in the PR body / return record for the exact one-line control per
guard and the sha256 before/after.
"""

import io
import zipfile
import shutil

import pytest

from app.utils.document_loader import (
    SheetExcelLoader,
    CELL_RANGE_LOCATOR_KEY,
    CELL_RANGE_LOCATOR_KIND,
    XLSX_HEADER_KEY,
    XLSX_HEADER_ROW_KEY,
)
from app.routes.document_routes import _UNIT_LOCATOR_KEYS

# Reuse the proven route driver + loader helpers from the capability suite.
from tests.utils.test_xlsx_capability import (  # noqa: F401 - pytest fixture `client`
    client,
    _embed,
    load_documents,
)

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# ===========================================================================
# SYNTHETIC fixture: two sheets, a merged title row, a vertical merged category,
# a header row, a cached formula, and a year cell.
# ===========================================================================


def make_SYNTHETIC_merged_workbook(path):
    """Revenue: merged title A1:C1, header row 2, data rows 3-4, cached SUM in B5.
    Detail:  header row 1, category "Hardware" vertically merged A2:A4, a yyyy year
    cell at C2. Ground truth is asserted directly by the tests below."""
    import datetime
    from openpyxl import Workbook

    wb = Workbook()

    rev = wb.active
    rev.title = "Revenue"
    rev["A1"] = "Quarterly Revenue"
    rev.merge_cells("A1:C1")               # horizontal merged title
    rev["A2"] = "Region"
    rev["B2"] = "Q1"
    rev["C2"] = "Q2"
    rev["A3"] = "EMEA"
    rev["B3"] = 100
    rev["C3"] = 110
    rev["A4"] = "AMER"
    rev["B4"] = 200
    rev["C4"] = 210
    rev["A5"] = "Total"
    rev["B5"] = "=SUM(B3:B4)"              # cached value injected below

    det = wb.create_sheet("Detail")
    det["A1"] = "Category"
    det["B1"] = "Item"
    det["C1"] = "Year"
    det["A2"] = "Hardware"
    det.merge_cells("A2:A4")               # vertical merged category label
    det["B2"] = "Widget"
    det["C2"] = datetime.datetime(2016, 1, 1)
    det["C2"].number_format = "yyyy"       # year-only date cell
    det["B3"] = "Gadget"
    det["C3"] = 2017
    det["B4"] = "Gizmo"
    det["C4"] = 2018

    wb.save(path)

    # Give the Revenue SUM the cached result Excel itself would have stored, so the
    # year-copy pass (written from cached values) keeps the total. Revenue is sheet1.
    tmp = str(path) + ".tmp"
    with zipfile.ZipFile(str(path)) as zin, zipfile.ZipFile(
        tmp, "w", zipfile.ZIP_DEFLATED
    ) as zout:
        for item in zin.infolist():
            payload = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                assert b"<f>SUM(B3:B4)</f>" in payload
                payload = payload.replace(
                    b"<f>SUM(B3:B4)</f>", b"<f>SUM(B3:B4)</f><v>300</v>"
                )
            zout.writestr(item, payload)
    shutil.move(tmp, str(path))


# ===========================================================================
# 1. Finer-than-sheet cell-range locator -- DETERMINISTIC (no parser dependency).
#    Control: delete the `doc.metadata[CELL_RANGE_LOCATOR_KEY] = ...` stamp in
#    _annotate, OR the cell_range computation in _sheet_locators -> this reddens.
# ===========================================================================


def test_SYNTHETIC_sheet_locators_expose_the_cell_range_extent(tmp_path):
    path = tmp_path / "merged.xlsx"
    make_SYNTHETIC_merged_workbook(str(path))

    locators = SheetExcelLoader(str(path))._sheet_locators()

    assert locators["Revenue"][CELL_RANGE_LOCATOR_KEY] == "Revenue!A1:C5"
    assert locators["Detail"][CELL_RANGE_LOCATOR_KEY] == "Detail!A1:C4"
    # Sheet-qualified so it is unique per unit even when two sheets share A1 notation.
    ranges = {v[CELL_RANGE_LOCATOR_KEY] for v in locators.values()}
    assert len(ranges) == 2


def test_SYNTHETIC_empty_sheet_has_no_cell_range(tmp_path):
    """A sheet with no content has no data extent to cite, so it is omitted rather
    than given a phantom range."""
    from openpyxl import Workbook

    path = tmp_path / "with-blank.xlsx"
    wb = Workbook()
    wb.active.title = "Data"
    wb.active["A1"] = "x"
    wb.active["B1"] = "y"
    wb.create_sheet("Blank")  # left empty
    wb.save(str(path))

    locators = SheetExcelLoader(str(path))._sheet_locators()
    assert "Blank" not in locators
    assert locators["Data"][CELL_RANGE_LOCATOR_KEY] == "Data!A1:B1"


# ===========================================================================
# 2. Header-row detection -- DETERMINISTIC.
#    Control: revert _detect_header_row's ">= 2 non-empty AND not all identical"
#    rule (e.g. accept the first non-empty row) -> the merged title is mis-picked
#    and header_row/header redden.
# ===========================================================================


def test_SYNTHETIC_header_row_detected_and_skips_merged_title(tmp_path):
    path = tmp_path / "merged.xlsx"
    make_SYNTHETIC_merged_workbook(str(path))

    locators = SheetExcelLoader(str(path))._sheet_locators()

    # Revenue: the merged title on row 1 must NOT be mistaken for the header.
    assert locators["Revenue"][XLSX_HEADER_ROW_KEY] == 2
    assert locators["Revenue"][XLSX_HEADER_KEY] == ["Region", "Q1", "Q2"]
    # Detail: the header is on row 1.
    assert locators["Detail"][XLSX_HEADER_ROW_KEY] == 1
    assert locators["Detail"][XLSX_HEADER_KEY] == ["Category", "Item", "Year"]


def test_detect_header_row_rule_is_first_multi_distinct_row():
    f = SheetExcelLoader._detect_header_row
    # A merged title (one anchor value) is skipped; the next multi-distinct row wins.
    assert f([(1, ["Title", None, None]), (2, ["Region", "Q1", "Q2"])]) == (
        2, ["Region", "Q1", "Q2"])
    # A filled merged title (all identical) is not a header.
    assert f([(1, ["H", "H", "H"]), (2, ["a", "b"])]) == (2, ["a", "b"])
    # A single-column list has no header.
    assert f([(1, ["only"]), (2, ["one"])]) == (None, None)
    # Empty input.
    assert f([]) == (None, None)


# ===========================================================================
# 3. Merged-cell resolution -- DETERMINISTIC via openpyxl on the produced copy.
#    Control: replace the anchor-fill `ws.cell(...).value = anchor` with `pass`
#    -> `found` never becomes True, _resolve_merged_cells returns None, and the
#    `copy_path is not None` assertion below reddens (phantom empties restored).
# ===========================================================================


def test_SYNTHETIC_merged_cells_are_unmerged_and_filled_no_phantom_empties(tmp_path):
    from openpyxl import load_workbook

    path = tmp_path / "merged.xlsx"
    make_SYNTHETIC_merged_workbook(str(path))

    copy_path = SheetExcelLoader(str(path))._resolve_merged_cells(
        str(path), str(tmp_path)
    )
    assert copy_path is not None, "a workbook with merged cells must produce a copy"

    wb = load_workbook(copy_path, data_only=True)
    try:
        for ws in wb.worksheets:
            assert list(ws.merged_cells.ranges) == [], (
                "%s still has merged ranges: %s" % (ws.title, ws.merged_cells.ranges))
        rev = wb["Revenue"]
        # Horizontal title: every cell of A1:C1 holds the title -> no phantom empties.
        assert rev["A1"].value == "Quarterly Revenue"
        assert rev["B1"].value == "Quarterly Revenue"
        assert rev["C1"].value == "Quarterly Revenue"
        det = wb["Detail"]
        # Vertical category: the label reaches every continuation row.
        assert det["A2"].value == "Hardware"
        assert det["A3"].value == "Hardware"
        assert det["A4"].value == "Hardware"
    finally:
        wb.close()


def test_resolve_merged_cells_is_a_noop_without_merges(tmp_path):
    """No merges -> parse the source untouched (None), the same contract the year
    pass uses; this is what keeps a plain workbook off the copy path."""
    from openpyxl import Workbook

    path = tmp_path / "plain.xlsx"
    wb = Workbook()
    wb.active["A1"] = "a"
    wb.active["B1"] = "b"
    wb.save(str(path))

    assert SheetExcelLoader(str(path))._resolve_merged_cells(
        str(path), str(tmp_path)) is None


def test_resolve_merged_cells_is_never_fatal_on_a_bad_workbook(tmp_path):
    """A diagnostic pass must never cost the content: an unreadable file returns
    None (parse the source as is), it does not raise."""
    path = tmp_path / "not.xlsx"
    path.write_bytes(b"PK\x03\x04 not a real package")
    assert SheetExcelLoader(str(path))._resolve_merged_cells(
        str(path), str(tmp_path)) is None


# ===========================================================================
# 4. The stamp reaches the LOADER Document, and the merged label reaches content.
#    (Parser-dependent: exercises UnstructuredExcelLoader. CODE-TESTED via image.)
#    Control: _resolve_merged_cells -> None makes "Hardware" appear once, reddening
#    the >= 3 assertion; deleting the _annotate stamp reddens the cell_range asserts.
# ===========================================================================


def test_SYNTHETIC_every_sheet_document_carries_cell_range_and_header(tmp_path):
    path = tmp_path / "merged.xlsx"
    make_SYNTHETIC_merged_workbook(str(path))

    docs = load_documents(path)
    expected_range = {"Revenue": "Revenue!A1:C5", "Detail": "Detail!A1:C4"}
    seen = set()
    for d in docs:
        sheet = d.metadata.get("page_name")
        assert d.metadata.get(CELL_RANGE_LOCATOR_KEY) == expected_range[sheet], (
            "%s chunk carries cell_range %r" % (sheet, d.metadata.get(
                CELL_RANGE_LOCATOR_KEY)))
        seen.add(sheet)
    assert seen == {"Revenue", "Detail"}
    # Header carried on the sheet's chunks.
    rev_headers = {
        tuple(d.metadata.get(XLSX_HEADER_KEY))
        for d in docs
        if d.metadata.get("page_name") == "Revenue" and d.metadata.get(XLSX_HEADER_KEY)
    }
    assert ("Region", "Q1", "Q2") in rev_headers


def test_SYNTHETIC_merged_category_label_reaches_every_row(tmp_path):
    """The anchor value of a vertical merge is not lost: after resolution the
    category appears once per spanned row rather than only on the first."""
    path = tmp_path / "merged.xlsx"
    make_SYNTHETIC_merged_workbook(str(path))

    docs = load_documents(path)
    detail = " ".join(
        d.page_content for d in docs if d.metadata.get("page_name") == "Detail"
    )
    assert detail.count("Hardware") >= 3, detail


# ===========================================================================
# 5. Receipt <-> chunk agreement for the new family, robust to placement.
#    Reads the EXPECTED kind from _UNIT_LOCATOR_KEYS so that when the FILES-ruled
#    placement lands, ONLY the asserted locator_kind value changes.
#    (Parser+route dependent. CODE-TESTED via image.)
# ===========================================================================


def _expected_kind_from_registry(stored_meta):
    for kind, key in _UNIT_LOCATOR_KEYS:
        if any((m or {}).get(key) is not None for m in stored_meta):
            return kind
    return "none"


def test_SYNTHETIC_cell_range_reaches_the_stored_chunk_and_receipt_agrees(
    client, tmp_path
):
    """The FILES-lead caution: prove the stamp reaches the STORED chunk, not only
    the loader Document; and the receipt names the family the store carries."""
    path = tmp_path / "merged.xlsx"
    make_SYNTHETIC_merged_workbook(str(path))

    r = _embed(client, "merged.xlsx", path.read_bytes(), file_id="f-cellrange")
    assert r.status_code == 200, r.text

    stored = [
        dict(d.metadata or {})
        for batch in client.inserted_batches
        for d in batch
    ]
    assert stored, "the workbook must store chunks"
    # Every stored chunk carries BOTH the sheet locator and the finer cell_range.
    assert all(m.get("page_name") for m in stored)
    assert all(m.get(CELL_RANGE_LOCATOR_KEY) for m in stored)

    receipt = r.json()["extraction"]
    assert receipt["locator_kind"] == _expected_kind_from_registry(stored)

    # CURRENT-STATE PIN: while cell_range is unregistered, XLSX still advertises
    # `sheet` (Option 1). When the placement lands this guard's branch retires and
    # the registry-driven assertion above carries the new value on its own.
    if (CELL_RANGE_LOCATOR_KIND, CELL_RANGE_LOCATOR_KEY) not in _UNIT_LOCATOR_KEYS:
        assert receipt["locator_kind"] == "sheet"
    # units_extracted stays one-per-sheet under either placement (cell_range is
    # sheet-qualified, so grouping by it yields the same count as by sheet).
    assert receipt["units_extracted"] == 2
