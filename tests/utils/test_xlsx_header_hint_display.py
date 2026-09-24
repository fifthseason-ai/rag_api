"""The XLSX `header` HINT shows what a cell DISPLAYS; the date itself is preserved.

F-XLSX-HINT-WORDING-AND-DATE-DISPLAY. Richard decision 4 (2026-09-23): preserve the
date, display the year.

WHY THIS EXISTS. The `header` hint is a convenience for a reader, so it has to read
like the sheet looks. A date cell formatted year-only (`yyyy`) is extracted as `2015`
in the TEXT (the year pass, F-XLSX-YEAR-FORMAT), so a hint that said
`2015-01-01 00:00:00` would disagree with the very text a citation points at — the
reader would see two different answers for one cell. The month and day are never lost:
the full date stays in `date_values`, keyed by sheet-qualified reference.

The hint mapping is deliberately narrow, and each boundary below is pinned with the
control that reddens it:
  * the fix        -> neutralise `_is_year_only_format(...)` in the header branch of
                      `_sheet_locators` (force `False`): TWO asserts red -- year display
                      AND hint-agrees-with-text (both depend on the mapping).
  * over-matching  -> force that predicate `True`: the full-date assert reds (a
                      `yyyy-mm-dd` cell must NOT be cut down to its year).
  * extent         -> STRUCTURAL, not mutation-proven, and deliberately labelled so. The
                      extent loop reads `cell.value` and never the mapped list
                      (document_loader.py:2105-2112), AND the year mapping is
                      emptiness-preserving (datetime -> int year, never None/""), while
                      the extent depends only on emptiness. So the extent is safe for two
                      independent reasons. An earlier draft of this docstring claimed
                      "feeding the mapped values to the extent scan reds `cell_range`";
                      reviewer RV-120 MEASURED that control (their M5) and it SURVIVES
                      86/86 -- exactly because of the emptiness-preserving property. Their
                      M6 (blank the mapped list AND make the extent read it) does red this
                      test. The test below pins the extent VALUE; it cannot, and no longer
                      claims to, detect an emptiness-preserving leak.
  * preservation   -> drop the `date_values` stamp: the preservation assert reds.

Hermetic: the loader is exercised directly. No database, no network, no embedding and
no rerank — nothing here makes a provider call.
"""
import datetime

import pytest

from app.utils.document_loader import (
    CELL_RANGE_LOCATOR_KEY,
    SheetExcelLoader,
    XLSX_HEADER_KEY,
    XLSX_HEADER_ROW_KEY,
)

# The header row deliberately mixes the five cases the mapping must tell apart:
# text, a year-only date, a full date, more text, and a number.
YEAR_ONLY = datetime.datetime(2015, 1, 1)      # displays as 2015
FULL_DATE = datetime.datetime(2016, 3, 4)      # displays as a full date


def make_header_hint_workbook(path):
    """One sheet 'Plan'. Row 1 is the header and carries, side by side: text, a
    YEAR-ONLY formatted date, a FULL-date formatted date, plain text, and a NUMBER. Row 2 is
    ordinary data so the header rule (>=2 non-empty, not all identical) picks row 1.

    A local builder rather than test_xlsx_year_cells.make_year_workbook: that one puts
    the year cell in a DATA row, which never reaches the header branch under test.
    """
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Plan"
    ws["A1"] = "Region"
    ws["B1"] = YEAR_ONLY
    ws["B1"].number_format = "yyyy"
    ws["C1"] = FULL_DATE
    ws["C1"].number_format = "yyyy-mm-dd"
    ws["D1"] = "Notes"
    ws["E1"] = 42          # a NUMERIC header cell (RV-120 N3: the name claimed one, the
                           # fixture had none -- a test naming a case it never built)
    ws["A2"] = "North"
    ws["B2"] = 10
    ws["C2"] = 20
    ws["D2"] = "ok"
    ws["E2"] = 7
    wb.save(path)
    return path


@pytest.fixture()
def hint_workbook(tmp_path):
    return str(make_header_hint_workbook(tmp_path / "plan.xlsx"))


def _locators(path):
    return SheetExcelLoader(path)._sheet_locators()


def test_header_hint_shows_a_year_only_date_cell_as_its_year(hint_workbook):
    """THE FIX. Without it the hint carries the raw datetime, which is exactly what
    the extracted text does NOT say."""
    loc = _locators(hint_workbook)["Plan"]
    assert loc[XLSX_HEADER_ROW_KEY] == 1, loc
    header = loc[XLSX_HEADER_KEY]
    assert header[1] == "2015", header
    assert "2015-01-01" not in header[1], header


def test_header_hint_does_not_reduce_a_full_date(hint_workbook):
    """OVER-MATCH GUARD. `yyyy-mm-dd` is not a year-only format; that cell keeps its
    whole date. A predicate that matched every date would red this."""
    header = _locators(hint_workbook)["Plan"][XLSX_HEADER_KEY]
    assert "2016-03-04" in header[2], header
    assert header[2] != "2016", header


def test_header_hint_leaves_text_and_numeric_cells_unchanged(hint_workbook):
    """NO REGRESSION: only year-only DATE cells are re-displayed."""
    header = _locators(hint_workbook)["Plan"][XLSX_HEADER_KEY]
    assert header[0] == "Region", header
    assert header[3] == "Notes", header
    assert header[4] == "42", header   # the NUMERIC header cell is untouched
    assert len(header) == 5, header


def test_cell_range_extent_is_unchanged_by_the_hint_display(hint_workbook):
    """The display mapping must not leak into extent detection, which reads the RAW
    value. The occupied extent is still the whole A1:E2 block.

    Pins the extent VALUE only. Per the header note, this cannot detect an
    emptiness-preserving leak (RV-120 M5 survives); RV-120 M6 is what reds it."""
    loc = _locators(hint_workbook)["Plan"]
    assert loc[CELL_RANGE_LOCATOR_KEY] == "Plan!A1:E2", loc


def test_the_full_date_is_preserved_in_date_values(hint_workbook):
    """Richard decision 4's other half: displaying the year must not destroy the date.
    The full date stays on the record, keyed sheet-qualified."""
    docs = SheetExcelLoader(hint_workbook).load()
    preserved = {}
    for doc in docs:
        preserved.update(doc.metadata.get("date_values") or {})
    assert preserved.get("Plan!B1") == "2015-01-01T00:00:00", preserved


def test_the_hint_agrees_with_the_extracted_text(hint_workbook):
    """The reason the fix exists: one cell must not read two ways. The hint says 2015
    and so does the text a citation lands on."""
    docs = SheetExcelLoader(hint_workbook).load()
    text = "\n".join(d.page_content for d in docs)
    header = _locators(hint_workbook)["Plan"][XLSX_HEADER_KEY]
    assert "2015" in text, text[:400]
    assert header[1] == "2015", header
    assert "2015-01-01 00:00:00" not in text, text[:400]
